from docx import Document as DocxDocument

from rag_lab import ingestion
from rag_lab.parsers import ParseResult, pick_parser
from rag_lab.parsers.docx import parse_docx


def _make_docx(path):
    doc = DocxDocument()
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph("Deep learning basics for the course.")
    doc.add_heading("2.3 Methods", level=2)
    doc.add_paragraph("We use gradient descent to train the model.")
    # Long enough that sentence chunking produces a chunk starting INSIDE the
    # Methods section (citations are anchored to the chunk start offset).
    for i in range(12):
        doc.add_paragraph(
            f"Additional paragraph {i} describing how backpropagation updates the "
            "weights layer by layer through the network."
        )
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Accuracy"
    table.cell(1, 1).text = "0.95"
    doc.save(str(path))


def test_docx_is_registered():
    assert pick_parser("x.docx") is parse_docx


def test_parse_docx_headings_become_markers(tmp_path):
    f = tmp_path / "report.docx"
    _make_docx(f)
    result = parse_docx(str(f))
    assert isinstance(result, ParseResult)
    assert "--- Introduction ---" in result.text
    assert "--- 2.3 Methods ---" in result.text
    assert "gradient descent" in result.text
    assert result.section_unit == "section"
    assert len(result.section_chars) == 2
    assert result.quality()["warnings"] == []


def test_parse_docx_heading_text_stays_inline(tmp_path):
    # Heading text must appear in the chunkable body too, not only in the
    # marker scaffolding, so chunks keep their section context.
    f = tmp_path / "report.docx"
    _make_docx(f)
    text = parse_docx(str(f)).text
    body_after_marker = text.split("--- 2.3 Methods ---", 1)[1]
    assert "2.3 Methods" in body_after_marker


def test_parse_docx_flattens_tables_row_per_line(tmp_path):
    f = tmp_path / "report.docx"
    _make_docx(f)
    text = parse_docx(str(f)).text
    assert "Metric | Value" in text
    assert "Accuracy | 0.95" in text


def test_parse_docx_empty_doc_warns(tmp_path):
    f = tmp_path / "empty.docx"
    DocxDocument().save(str(f))
    q = parse_docx(str(f)).quality()
    assert q["total_chars"] == 0
    assert any("no text extracted" in w for w in q["warnings"])


def test_docx_citations_use_section_labels(tmp_path):
    f = tmp_path / "report.docx"
    _make_docx(f)
    parsed = parse_docx(str(f))
    chunks = ingestion.make_chunks(parsed.effective_text, "sentence", 512, 64)
    metas = ingestion.build_metadatas(
        str(f), parsed.effective_text, chunks, "sha1", "sentence", 512, 64
    )
    citations = [m["citation"] for m in metas]
    assert any(c == "report.docx §2.3 Methods" for c in citations)


def test_marker_regex_allows_hyphenated_headings():
    text = "\n\n--- State-of-the-art Models ---\n\nTransformers dominate."
    markers = ingestion._markers(text)
    assert markers and markers[0][1] == "State-of-the-art Models"


def test_marker_regex_page_priority_unchanged():
    text = "\n\n--- Page 12 ---\n\nSome page text."
    markers = ingestion._markers(text)
    assert markers[0][1] == "Page 12"
    assert ingestion._citation("a.pdf", markers[0][1], 0) == "a.pdf p.12"
