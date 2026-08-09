from pptx import Presentation as PptxPresentation

from rag_lab import ingestion
from rag_lab.parsers import ParseResult, pick_parser
from rag_lab.parsers.pptx import parse_pptx


def _make_pptx(path, with_notes=True):
    prs = PptxPresentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])  # Titel + Inhalt
    slide.shapes.title.text = "KI im Finanzwesen"
    slide.placeholders[1].text = "Vorhersage von Kreditausfällen\nBetrugserkennung"
    slide2 = prs.slides.add_slide(prs.slide_layouts[5])  # nur Titel
    slide2.shapes.title.text = "Kernmodell"
    if with_notes:
        slide2.notes_slide.notes_text_frame.text = (
            "Kreditrisikobewertung per logistischer Regression"
        )
    prs.save(str(path))


def test_pptx_is_registered():
    assert pick_parser("lecture03.pptx") is parse_pptx


def test_parse_pptx_slides_become_markers(tmp_path):
    f = tmp_path / "lecture03.pptx"
    _make_pptx(f)
    result = parse_pptx(str(f))
    assert isinstance(result, ParseResult)
    assert "--- Slide 1 ---" in result.text
    assert "--- Slide 2 ---" in result.text
    assert "KI im Finanzwesen" in result.text
    assert result.section_unit == "slide"
    assert len(result.section_chars) == 2


def test_parse_pptx_includes_speaker_notes(tmp_path):
    f = tmp_path / "lecture03.pptx"
    _make_pptx(f, with_notes=True)
    text = parse_pptx(str(f)).text
    assert "Notes: Kreditrisikobewertung per logistischer Regression" in text


def test_parse_pptx_without_notes_has_no_notes_marker(tmp_path):
    f = tmp_path / "lecture03.pptx"
    _make_pptx(f, with_notes=False)
    text = parse_pptx(str(f)).text
    assert "Notes:" not in text


def test_pptx_citations_use_slide_labels(tmp_path):
    f = tmp_path / "lecture03.pptx"
    _make_pptx(f)
    assert ingestion._citation("lecture03.pptx", "Slide 12", 0) == "lecture03.pptx slide 12"
