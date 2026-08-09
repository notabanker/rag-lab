"""PPTX parser: each slide becomes a section; speaker notes join the slide body."""
from pptx import Presentation

from .base import ParseResult


def _slide_lines(slide) -> list[str]:
    lines = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                text = " ".join(para.text.split())
                if text:
                    lines.append(text)
        if shape.has_table:
            for row in shape.table.rows:
                cells = [" ".join(c.text.split()) for c in row.cells]
                line = " | ".join(c for c in cells if c)
                if line:
                    lines.append(line)
    return lines


def parse_pptx(path: str) -> ParseResult:
    prs = Presentation(path)
    parts = []
    section_chars = []
    for i, slide in enumerate(prs.slides, 1):
        body = "\n".join(_slide_lines(slide))
        if slide.has_notes_slide:
            notes = " ".join(slide.notes_slide.notes_text_frame.text.split())
            if notes:
                body += f"\nNotes: {notes}"
        section_chars.append(len(body.strip()))
        parts.append(f"\n\n--- Slide {i} ---\n\n{body}")
    return ParseResult(text="".join(parts), section_unit="slide", section_chars=section_chars)
