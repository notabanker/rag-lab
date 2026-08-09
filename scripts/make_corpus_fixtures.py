"""Generate the corpus fixtures for the golden eval set (data/docx, data/pptx).

Idempotent: rerun to regenerate. Requires the main dependency group
(python-docx, python-pptx).
"""
from pathlib import Path

from docx import Document
from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent / "data"


def _ml_notes() -> Document:
    doc = Document()
    doc.add_heading("Maschinelles Lernen — Einführung", level=1)
    doc.add_paragraph(
        "Unter überwachtem Lernen versteht man das Training mit gelabelten Daten, "
        "bei dem das Modell von Eingaben auf bekannte Ausgaben lernt."
    )
    doc.add_paragraph(
        "Beim unüberwachten Lernen findet das Modell selbst Strukturen in den Daten, "
        "ohne vorgegebene Labels."
    )
    doc.add_heading("Lineare Regression", level=1)
    doc.add_paragraph(
        "Die lineare Regression modelliert den Zusammenhang y = w·x + b zwischen "
        "einer Eingabevariablen und einer Zielgröße."
    )
    doc.add_paragraph(
        "Der Gradientenabstieg minimiert die Verlustfunktion, indem er die Gewichte "
        "iterativ in Richtung des steilsten Abstiegs anpasst."
    )
    doc.add_heading("Evaluationsmetriken", level=1)
    doc.add_paragraph("Der F1-Score ist das harmonische Mittel aus Precision und Recall.")
    return doc


def _thesis_guide() -> Document:
    doc = Document()
    doc.add_heading("Formale Anforderungen", level=1)
    doc.add_paragraph("Die Bachelorarbeit umfasst in der Regel 60 Seiten Text.")
    doc.add_paragraph("Das Manuskript wird in 12-Punkt-Schrift gesetzt.")
    doc.add_heading("Abgabe und Fristen", level=1)
    doc.add_paragraph("Die Abgabefrist endet am 15. September jedes Jahres.")
    return doc


def _lecture03() -> Presentation:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "KI im Finanzwesen"
    slide.placeholders[1].text = "Vorhersage von Kreditausfällen\nBetrugserkennung in Echtzeit"
    slide2 = prs.slides.add_slide(prs.slide_layouts[1])
    slide2.shapes.title.text = "Anwendungen"
    slide2.placeholders[1].text = (
        "Robo-Advisory automatisiert die Anlageberatung\n"
        "Chatbots beantworten Kundenanfragen rund um die Uhr"
    )
    slide3 = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only (Blank [6] hat keinen Titel-Placeholder)
    slide3.shapes.title.text = "Kernmodell"
    slide3.notes_slide.notes_text_frame.text = (
        "Das Kernmodell ist die Kreditrisikobewertung per logistischer Regression."
    )
    return prs


def main():
    (ROOT / "docx").mkdir(parents=True, exist_ok=True)
    (ROOT / "pptx").mkdir(parents=True, exist_ok=True)
    _ml_notes().save(ROOT / "docx" / "machine_learning_notes.docx")
    _thesis_guide().save(ROOT / "docx" / "bachelor_thesis_guide.docx")
    _lecture03().save(ROOT / "pptx" / "lecture03.pptx")
    print("Fixtures written to data/docx and data/pptx")


if __name__ == "__main__":
    main()
